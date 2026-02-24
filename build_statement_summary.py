#!/usr/bin/env python3
"""Build an Excel summary of statement inflows/outflows from screenshot data."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


def parse_de_amount(value: str) -> Decimal:
    """Convert German-formatted amount like 18.833,33 into Decimal."""
    normalized = value.replace(".", "").replace(",", ".")
    return Decimal(normalized)


def format_sheet(sheet, title: str, rows: list[tuple[str, Decimal]]) -> Decimal:
    """Write grouped rows into worksheet and return total."""
    sheet.title = title
    sheet["A1"] = "Наименование (DE)"
    sheet["B1"] = "Сумма"
    sheet["A1"].font = Font(bold=True)
    sheet["B1"].font = Font(bold=True)

    total = Decimal("0")
    current_row = 2
    for name, amount in rows:
        sheet.cell(row=current_row, column=1, value=name)
        amount_cell = sheet.cell(row=current_row, column=2, value=float(amount))
        amount_cell.number_format = '#,##0.00'
        total += amount
        current_row += 1

    sheet.cell(row=current_row, column=1, value="ИТОГО").font = Font(bold=True)
    total_cell = sheet.cell(row=current_row, column=2, value=float(total))
    total_cell.font = Font(bold=True)
    total_cell.number_format = '#,##0.00'

    sheet.column_dimensions["A"].width = 40
    sheet.column_dimensions["B"].width = 16
    return total


def main() -> None:
    # Transactions read from the screenshot supplied in chat.
    # H = income (приход), S = expense (расход).
    transactions = [
        ("Umbuchung", "30,00", "S"),
        ("Überweisungsauftrag", "250,00", "S"),
        ("Dauerauftragsgutschr", "38,35", "H"),
        ("Überweisungsgutschr.", "19.000,00", "H"),
        ("Basislastschrift", "18.833,33", "S"),
        ("Abschluss lt. Anlage 1", "12,84", "S"),
        ("Überweisungsgutschr.", "30,00", "H"),
        ("Abschluss lt. Anlage 1", "11,48", "S"),
    ]

    income: dict[str, Decimal] = {}
    expense: dict[str, Decimal] = {}

    for tx_name, tx_amount, tx_type in transactions:
        amount = parse_de_amount(tx_amount)
        target = income if tx_type == "H" else expense
        target[tx_name] = target.get(tx_name, Decimal("0")) + amount

    wb = Workbook()
    income_ws = wb.active
    income_total = format_sheet(income_ws, "Приходы", list(income.items()))
    expense_ws = wb.create_sheet("Расходы")
    expense_total = format_sheet(expense_ws, "Расходы", list(expense.items()))

    summary_ws = wb.create_sheet("Сводка")
    summary_ws["A1"] = "Показатель"
    summary_ws["B1"] = "Сумма"
    summary_ws["A1"].font = Font(bold=True)
    summary_ws["B1"].font = Font(bold=True)
    summary_ws["A2"] = "Общий приход"
    summary_ws["B2"] = float(income_total)
    summary_ws["A3"] = "Общий расход"
    summary_ws["B3"] = float(expense_total)
    summary_ws["A4"] = "Сальдо (приход - расход)"
    summary_ws["B4"] = float(income_total - expense_total)
    for row_idx in (2, 3, 4):
        summary_ws[f"B{row_idx}"].number_format = '#,##0.00'
    summary_ws.column_dimensions["A"].width = 34
    summary_ws.column_dimensions["B"].width = 16

    output_path = Path("statement_summary.xlsx")
    wb.save(output_path)
    print(f"Created {output_path.resolve()}")


if __name__ == "__main__":
    main()
