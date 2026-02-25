#!/usr/bin/env python3
"""Build income/expense summary workbook from OCR-ed PDF transactions."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


def parse_de_amount(value: str) -> Decimal:
    """Convert German amount format (1.234,56) to Decimal."""
    return Decimal(value.replace(".", "").replace(",", "."))


def write_grouped_sheet(ws, sheet_name: str, grouped: dict[str, Decimal]) -> Decimal:
    ws.title = sheet_name
    ws["A1"] = "Наименование (DE)"
    ws["B1"] = "Сумма"
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)

    total = Decimal("0")
    row = 2
    for tx_name, amount in grouped.items():
        ws.cell(row=row, column=1, value=tx_name)
        amount_cell = ws.cell(row=row, column=2, value=float(amount))
        amount_cell.number_format = '#,##0.00'
        total += amount
        row += 1

    ws.cell(row=row, column=1, value="ИТОГО").font = Font(bold=True)
    total_cell = ws.cell(row=row, column=2, value=float(total))
    total_cell.number_format = '#,##0.00'
    total_cell.font = Font(bold=True)

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 16
    return total


def main() -> None:
    # Transactions transcribed from OCR of:
    # /home/ubuntu/.cursor/projects/workspace/uploads/Ko_ln-25022026102427.pdf
    # H = income, S = expense.
    transactions = [
        ("Dauerauftragsgutschr", "1.650,00", "H"),
        ("Dauerauftragsgutschr", "400,60", "H"),
        ("Überweisungsgutschr.", "800,00", "H"),
        ("Dauerauftragsgutschr", "750,00", "H"),
        ("Dauerauftragsgutschr", "1.200,00", "H"),
        ("Überweisungsgutschr.", "1.750,00", "H"),
        ("Basislastschrift", "70,00", "S"),
        ("Basislastschrift", "91,00", "S"),
        ("Basislastschrift", "362,00", "S"),
        ("Basislastschrift", "77,00", "S"),
        ("Überweisungsgutschr.", "950,00", "H"),
        ("Basislastschrift", "47,90", "S"),
        ("Überweisungsauftrag", "825,86", "S"),
        ("Überweisungsauftrag", "5.644,17", "S"),
        ("Überweisungsauftrag", "827,18", "S"),
        ("Überweisungsauftrag", "1.047,20", "S"),
        ("Überweisungsauftrag", "244,37", "S"),
        ("Überweisungsauftrag", "7,12", "S"),
        ("Überweisungsauftrag", "496,06", "S"),
        ("Basislastschrift", "50,90", "S"),
        ("Überweisungsgutschr.", "115,68", "H"),
        ("Überweisungsauftrag", "2.303,29", "S"),
        ("Überweisungsgutschr.", "44,54", "H"),
        ("Basislastschrift", "41,95", "S"),
        ("Basislastschrift", "69,88", "S"),
        ("Überweisungsgutschr.", "417,69", "H"),
        ("Überweisungsgutschr.", "236,53", "H"),
        ("Überweisungsgutschr.", "1.170,00", "H"),
        ("Basislastschrift", "186,52", "S"),
        ("Dauerauftragsgutschr", "950,00", "H"),
        ("Dauerauftragsgutschr", "470,00", "H"),
        ("Überweisungsauftrag", "11,65", "S"),
        ("Basislastschrift", "38,09", "S"),
        ("Überweisungsauftrag", "27,37", "S"),
        ("Überweisungsauftrag", "15,47", "S"),
        ("Dauerauftragsgutschr", "1.300,00", "H"),
        ("Basislastschrift", "63,00", "S"),
        ("Basislastschrift", "149,00", "S"),
        ("Überweisungsgutschr.", "1.050,00", "H"),
        ("Überweisungsgutschr.", "1.400,00", "H"),
        ("Basislastschrift", "10.124,99", "S"),
        ("Dauerauftragsgutschr", "890,00", "H"),
        ("Überweisungsgutschr.", "1.350,00", "H"),
        ("Abschluss lt. Anlage 1", "28,76", "S"),
    ]

    income_by_type: dict[str, Decimal] = {}
    expense_by_type: dict[str, Decimal] = {}

    for tx_type, amount_str, marker in transactions:
        amount = parse_de_amount(amount_str)
        bucket = income_by_type if marker == "H" else expense_by_type
        bucket[tx_type] = bucket.get(tx_type, Decimal("0")) + amount

    wb = Workbook()
    incomes_ws = wb.active
    income_total = write_grouped_sheet(incomes_ws, "Доходы", income_by_type)
    expenses_ws = wb.create_sheet("Расходы")
    expense_total = write_grouped_sheet(expenses_ws, "Расходы", expense_by_type)

    summary_ws = wb.create_sheet("Сводка")
    summary_ws["A1"] = "Показатель"
    summary_ws["B1"] = "Сумма"
    summary_ws["A1"].font = Font(bold=True)
    summary_ws["B1"].font = Font(bold=True)
    summary_ws["A2"] = "Общий доход (H)"
    summary_ws["B2"] = float(income_total)
    summary_ws["A3"] = "Общий расход (S)"
    summary_ws["B3"] = float(expense_total)
    summary_ws["A4"] = "Сальдо (доход - расход)"
    summary_ws["B4"] = float(income_total - expense_total)
    for idx in (2, 3, 4):
        summary_ws[f"B{idx}"].number_format = '#,##0.00'
    summary_ws.column_dimensions["A"].width = 30
    summary_ws.column_dimensions["B"].width = 16

    output_path = Path("/workspace/testsummuryfrompdf.xlsx")
    wb.save(output_path)
    print(f"Created {output_path}")
    print(f"Income total: {income_total}")
    print(f"Expense total: {expense_total}")


if __name__ == "__main__":
    main()
